"""REPORT.md → REPORT.docx（python-docx，处理本报告用到的 Markdown 子集）

用法: python build_report_docx.py [SRC.md] [DST.docx]   （默认 REPORT.md → REPORT.docx）
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).parent
SRC = ROOT / (sys.argv[1] if len(sys.argv) > 1 else "REPORT.md")
DST = ROOT / (sys.argv[2] if len(sys.argv) > 2 else "REPORT.docx")

CJK = "微软雅黑"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
GRAY = RGBColor(0x66, 0x66, 0x66)


def set_font(run, size=10.5, bold=False, color=None, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_inline(par, text, size=10.5, base_color=None, italic=False):
    """Parse **bold** and `code` inline markup."""
    text = text.replace("`", "")
    for part in re.split(r"(\*\*.+?\*\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            set_font(par.add_run(part[2:-2]), size=size, bold=True,
                     color=base_color, italic=italic)
        else:
            set_font(par.add_run(part), size=size, color=base_color, italic=italic)


def strip_links(s: str) -> str:
    return re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", s)


def main():
    lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()
    # page margins
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(0.9)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        if not line.strip() or line.strip() == "---":
            i += 1
            continue

        # headings
        m = re.match(r"^(#{1,3})\s+(.*)", line)
        if m:
            level = len(m.group(1))
            text = strip_links(m.group(2))
            h = doc.add_heading("", level=level)
            run = h.add_run(text)
            set_font(run, size={1: 17, 2: 14, 3: 12}[level], bold=True, color=ACCENT)
            i += 1
            continue

        # image
        m = re.match(r"^!\[[^\]]*\]\(([^)]+)\)", line.strip())
        if m:
            img = ROOT / m.group(1)
            if img.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(img), width=Inches(6.3))
            i += 1
            continue

        # table block
        if line.strip().startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            rows = [
                [c.strip() for c in r.strip("|").split("|")]
                for r in block
                if not re.match(r"^\|[\s:|-]+\|$", r)
            ]
            if not rows:
                continue
            tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
            tbl.style = "Light Grid Accent 1"
            for ri, row in enumerate(rows):
                for ci, cell in enumerate(row):
                    if ci >= len(tbl.rows[ri].cells):
                        continue
                    par = tbl.rows[ri].cells[ci].paragraphs[0]
                    add_inline(par, strip_links(cell), size=9.5,
                               base_color=None if ri else ACCENT)
                    if ri == 0:
                        for r_ in par.runs:
                            r_.bold = True
            doc.add_paragraph()
            continue

        # blockquote
        if line.strip().startswith(">"):
            text = strip_links(line.strip().lstrip("> ").strip())
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_inline(p, text, size=9.5, base_color=GRAY, italic=True)
            i += 1
            continue

        # lists
        m = re.match(r"^(\d+)\.\s+(.*)", line.strip())
        if m:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, strip_links(m.group(2)))
            i += 1
            continue
        if line.strip().startswith(("- ", "* ")):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, strip_links(line.strip()[2:]))
            i += 1
            continue

        # footer italic line (*...*)
        s = line.strip()
        if s.startswith("*") and s.endswith("*") and not s.startswith("**"):
            p = doc.add_paragraph()
            add_inline(p, strip_links(s.strip("*")), size=9, base_color=GRAY, italic=True)
            i += 1
            continue

        # normal paragraph
        p = doc.add_paragraph()
        add_inline(p, strip_links(line.strip()))
        i += 1

    doc.save(DST)
    print(f"Saved {DST} ({DST.stat().st_size/1024:.0f} KB)")


if __name__ == "__main__":
    main()
